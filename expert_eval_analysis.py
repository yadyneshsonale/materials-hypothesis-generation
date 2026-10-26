"""Experiment 30 (loop.md section 13.1): analyse a returned expert rating sheet against the AI judges.

Usage:
    python expert_eval_analysis.py expert_evaluation/5_returned_ratings/expert_ratings_AB.xlsx
    python expert_eval_analysis.py path/to/csv_folder        # the four CSV files instead of the workbook

Pre-registered analysis (written before any rating existed):
  primary    Spearman rho between the expert's agreement (mean of intervention, conditions, mechanism, outcome; 0-4)
             and the GPT-5.6 reconstruction score over the generated items (o3-mini writer and Llama baseline), with a
             question-clustered bootstrap 95% CI. Bands: >= 0.6 tracks the expert, 0.3-0.6 partial, < 0.3 unreliable.
  secondary  per-dimension and DeepSeek rho; rho over all 60 items; within-question order agreement; expert-rated
             o3-mini minus Llama (agreement and Part-1 quality); the reference ceiling; answer-key accuracy by
             extractor; question sense and hints; identification of the reference against chance (1/3).
Only complete ratings are used, so a partly completed sheet (for example Q01-Q10) is analysed as is.
Outputs: experiments/.../expert_eval/analysis_<rater>.json and .md
"""

from __future__ import annotations

import argparse
import csv
import json
import math
import random
import statistics
from pathlib import Path
from typing import Any

from hypothesis_evidence_v2 import atomic_json

ROOT = Path(__file__).resolve().parent
INTERNAL = ROOT / "experiments" / "high_temperature_hea_qwen35b" / "expert_eval"
DIMS = ("intervention_reconstruction", "condition_agreement", "mechanism_agreement", "outcome_agreement")
SHEET_COLUMNS = {
    "Part 1 – hypotheses": ("plausibility", "testability", "usefulness"),
    "Part 1 – questions": ("question_sense", "closest"),
    "Part 2 – hypotheses": ("intervention", "conditions", "mechanism", "outcome"),
    "Part 2 – questions": ("key_accuracy", "hint"),
}
CSV_NAMES = {"Part 1 – hypotheses": "part_1_hypotheses.csv", "Part 1 – questions": "part_1_questions.csv",
             "Part 2 – hypotheses": "part_2_hypotheses.csv", "Part 2 – questions": "part_2_questions.csv"}
GENERATED = ("o3mini_writer", "llama_baseline")
BOOTSTRAP = 10_000
SEED = 20261026


# ---------------------------------------------------------------- reading

def _value(raw: Any) -> Any:
    if raw is None:
        return None
    text = str(raw).strip()
    if not text:
        return None
    try:
        number = float(text)
        return int(number) if number.is_integer() else number
    except ValueError:
        return text


def read_ratings(path: Path) -> dict[str, Any]:
    """Return {sheet: {key: {field: value}}}, key = 'Q01-A' for hypothesis sheets and 'Q01' for question sheets."""
    tables: dict[str, list[list[Any]]] = {}
    rater = path.stem.replace("expert_ratings_", "") or "expert"
    if path.is_dir():
        for sheet, name in CSV_NAMES.items():
            with open(path / name, encoding="utf-8") as handle:
                tables[sheet] = list(csv.reader(handle))
    else:
        from openpyxl import load_workbook
        wb = load_workbook(path, data_only=True)
        for sheet in SHEET_COLUMNS:
            tables[sheet] = [list(row) for row in wb[sheet].iter_rows(values_only=True)]
        name = _value(wb["START HERE"]["B3"].value)
        rater = str(name) if name else rater
    out: dict[str, Any] = {"rater": rater}
    for sheet, fields in SHEET_COLUMNS.items():
        rows = tables[sheet][1:]
        per_hypothesis = "hypotheses" in sheet
        data = {}
        for row in rows:
            if not row or _value(row[0]) is None:
                continue
            key = f"{row[0]}-{row[1]}" if per_hypothesis else str(row[0])
            start = 2 if per_hypothesis else 1
            values = [_value(row[start + i]) if start + i < len(row) else None for i in range(len(fields))]
            data[key] = dict(zip(fields, values))
            data[key]["comment"] = _value(row[start + len(fields)]) if start + len(fields) < len(row) else None
        out[sheet] = data
    return out


# ---------------------------------------------------------------- statistics

def _ranks(values: list[float]) -> list[float]:
    order = sorted(range(len(values)), key=lambda i: values[i])
    ranks = [0.0] * len(values)
    i = 0
    while i < len(order):
        j = i
        while j + 1 < len(order) and values[order[j + 1]] == values[order[i]]:
            j += 1
        for k in range(i, j + 1):
            ranks[order[k]] = (i + j) / 2 + 1
        i = j + 1
    return ranks


def spearman(x: list[float], y: list[float]) -> float | None:
    if len(x) < 3 or len(set(x)) < 2 or len(set(y)) < 2:
        return None
    rx, ry = _ranks(x), _ranks(y)
    mx, my = statistics.fmean(rx), statistics.fmean(ry)
    num = sum((a - mx) * (b - my) for a, b in zip(rx, ry))
    den = math.sqrt(sum((a - mx) ** 2 for a in rx) * sum((b - my) ** 2 for b in ry))
    return num / den


def clustered_rho(pairs: list[tuple[str, float, float]]) -> dict[str, Any]:
    """Spearman rho over (question, x, y) items, with a question-clustered bootstrap 95% CI."""
    rho = spearman([p[1] for p in pairs], [p[2] for p in pairs])
    if rho is None:
        return {"n_items": len(pairs), "rho": None}
    by_q: dict[str, list[tuple[str, float, float]]] = {}
    for p in pairs:
        by_q.setdefault(p[0], []).append(p)
    qs = sorted(by_q)
    rng = random.Random(SEED)
    boots = []
    for _ in range(BOOTSTRAP):
        sample = [item for q in (rng.choice(qs) for _ in qs) for item in by_q[q]]
        value = spearman([p[1] for p in sample], [p[2] for p in sample])
        if value is not None:
            boots.append(value)
    boots.sort()
    low, high = boots[int(0.025 * len(boots))], boots[min(len(boots) - 1, int(0.975 * len(boots)))]
    return {"n_items": len(pairs), "n_questions": len(qs), "rho": round(rho, 3), "ci95": [round(low, 3), round(high, 3)]}


def paired_by_question(a: dict[str, float], b: dict[str, float]) -> dict[str, Any]:
    """Mean of (a - b) over questions with a clustered bootstrap CI and an exact sign test."""
    keys = sorted(set(a) & set(b))
    diffs = [a[k] - b[k] for k in keys]
    if not diffs:
        return {"n_questions": 0}
    rng = random.Random(SEED)
    boots = sorted(statistics.fmean(rng.choice(diffs) for _ in diffs) for _ in range(BOOTSTRAP))
    up, down = sum(d > 0 for d in diffs), sum(d < 0 for d in diffs)
    n = up + down
    tail = sum(math.comb(n, k) for k in range(0, min(up, down) + 1)) / 2 ** n if n else 1.0
    return {"n_questions": len(keys), "diff": round(statistics.fmean(diffs), 3),
            "ci95": [round(boots[int(0.025 * BOOTSTRAP)], 3), round(boots[int(0.975 * BOOTSTRAP) - 1], 3)],
            "up": up, "down": down, "sign_p": round(min(1.0, 2 * tail), 4)}


def binomial_greater(k: int, n: int, p: float) -> float:
    return sum(math.comb(n, i) * p ** i * (1 - p) ** (n - i) for i in range(k, n + 1))


def band(rho: float | None) -> str:
    if rho is None:
        return "not computable"
    return "tracks the expert" if rho >= 0.6 else "partial" if rho >= 0.3 else "unreliable"


# ---------------------------------------------------------------- analysis

def analyse(ratings: dict[str, Any], key: dict[str, Any], scores: dict[str, Any]) -> dict[str, Any]:
    source = {f"{q['qid']}-{lab}": q["labels"][lab]["source"] for q in key["questions"] for lab in "ABC"}
    qset = {q["qid"]: q["set"] for q in key["questions"]}
    reference_label = {q["qid"]: next(lab for lab in "ABC" if q["labels"][lab]["source"] == "reference")
                       for q in key["questions"]}
    p2 = {k: v for k, v in ratings["Part 2 – hypotheses"].items()
          if all(isinstance(v[f], (int, float)) for f in ("intervention", "conditions", "mechanism", "outcome"))}
    expert_agree = {k: statistics.fmean(v[f] for f in ("intervention", "conditions", "mechanism", "outcome"))
                    for k, v in p2.items()}
    p1 = {k: v for k, v in ratings["Part 1 – hypotheses"].items()
          if all(isinstance(v[f], (int, float)) for f in ("plausibility", "testability", "usefulness"))}
    out: dict[str, Any] = {"rater": ratings["rater"],
                           "completed": {"part1_items": len(p1), "part2_items": len(p2)}}

    def pairs(items: list[str], judge: str, expert: dict[str, float], judge_field: str = "recon",
              dim: str | None = None) -> list[tuple[str, float, float]]:
        res = []
        for k in items:
            if k in expert and k in scores:
                j = scores[k][judge]["scores"][dim] if dim else scores[k][judge][judge_field]
                res.append((k.split("-")[0], expert[k], j))
        return res

    generated = sorted(k for k in expert_agree if source[k] in GENERATED)
    everything = sorted(expert_agree)
    out["primary"] = {**clustered_rho(pairs(generated, "gpt56", expert_agree)), "scope": "40 generated items"}
    out["primary"]["interpretation"] = band(out["primary"].get("rho"))
    dim_fields = {"intervention_reconstruction": "intervention", "condition_agreement": "conditions",
                  "mechanism_agreement": "mechanism", "outcome_agreement": "outcome"}
    out["per_dimension_gpt56"] = {
        dim: clustered_rho(pairs(generated, "gpt56", {k: float(p2[k][field]) for k in p2}, dim=dim))
        for dim, field in dim_fields.items()}
    out["deepseek"] = clustered_rho(pairs(generated, "deepseek", expert_agree))
    out["all_60_gpt56"] = clustered_rho(pairs(everything, "gpt56", expert_agree))
    # Within-question order of the two generated hypotheses.
    agree = disagree = 0
    for q in sorted({k.split("-")[0] for k in generated}):
        ks = {source[k]: k for k in generated if k.startswith(q + "-")}
        if len(ks) < 2:
            continue
        e = expert_agree[ks["o3mini_writer"]] - expert_agree[ks["llama_baseline"]]
        j = scores[ks["o3mini_writer"]]["gpt56"]["recon"] - scores[ks["llama_baseline"]]["gpt56"]["recon"]
        if e == 0 or j == 0:
            continue
        agree += (e > 0) == (j > 0)
        disagree += (e > 0) != (j > 0)
    out["order_agreement_gpt56"] = {"questions": agree + disagree,
                                    "agree_share": round(agree / (agree + disagree), 3) if agree + disagree else None}

    def per_question(values: dict[str, float], src: str) -> dict[str, float]:
        return {k.split("-")[0]: v for k, v in values.items() if source[k] == src}

    out["expert_o3mini_minus_llama"] = {"agreement": paired_by_question(per_question(expert_agree, "o3mini_writer"),
                                                                        per_question(expert_agree, "llama_baseline"))}
    for field in ("plausibility", "testability", "usefulness"):
        vals = {k: float(v[field]) for k, v in p1.items()}
        out["expert_o3mini_minus_llama"][field] = paired_by_question(per_question(vals, "o3mini_writer"),
                                                                     per_question(vals, "llama_baseline"))
    out["by_source"] = {}
    for src in ("o3mini_writer", "llama_baseline", "reference"):
        ks = [k for k in source if source[k] == src]
        out["by_source"][src] = {
            "expert_agreement": _mean([expert_agree[k] for k in ks if k in expert_agree]),
            "gpt56_recon": _mean([scores[k]["gpt56"]["recon"] for k in ks if k in expert_agree]),
            "deepseek_recon": _mean([scores[k]["deepseek"]["recon"] for k in ks if k in expert_agree]),
            **{field: _mean([p1[k][field] for k in ks if k in p1]) for field in ("plausibility", "testability",
                                                                                 "usefulness")}}
    pq2 = ratings["Part 2 – questions"]
    out["answer_key_accuracy"] = {
        name: _mean([pq2[q]["key_accuracy"] for q in pq2 if isinstance(pq2[q]["key_accuracy"], (int, float))
                     and (qset.get(q) in sets)])
        for name, sets in (("experiment3_keys (original 12)", {"original12"}),
                           ("gpt56_keys (test2, test3)", {"test2", "test3"}))}
    hints = [pq2[q]["hint"] for q in pq2 if pq2[q]["hint"] in ("No", "Slightly", "Strongly")]
    out["question_hints"] = {h: hints.count(h) for h in ("No", "Slightly", "Strongly")}
    pq1 = ratings["Part 1 – questions"]
    out["question_sense"] = _mean([v["question_sense"] for v in pq1.values()
                                   if isinstance(v["question_sense"], (int, float))])
    guesses = {q: v["closest"] for q, v in pq1.items() if v["closest"] in ("A", "B", "C")}
    correct = sum(guesses[q] == reference_label.get(q) for q in guesses)
    out["reference_identification"] = {
        "guesses": len(guesses), "correct": correct,
        "cant_tell": sum(1 for v in pq1.values() if v["closest"] == "Can't tell"),
        "rate": round(correct / len(guesses), 3) if guesses else None,
        "p_vs_chance_one_third": round(binomial_greater(correct, len(guesses), 1 / 3), 4) if guesses else None}
    return out


def _mean(values: list[Any]) -> float | None:
    values = [v for v in values if isinstance(v, (int, float))]
    return round(statistics.fmean(values), 3) if values else None


def render(result: dict[str, Any]) -> str:
    p = result["primary"]
    lines = [f"# Expert evaluation: analysis for rater {result['rater']}", "",
             f"Completed: {result['completed']['part1_items']} Part-1 items, {result['completed']['part2_items']} "
             "Part-2 items.", "",
             f"**Primary (judge validity):** Spearman ρ between expert agreement and GPT-5.6 reconstruction over "
             f"{p.get('n_items')} generated items = **{p.get('rho')}** {p.get('ci95', '')} → *{p['interpretation']}*.",
             "", "| Measure | ρ | 95% CI | items |", "|---|---:|---|---:|"]
    for name, r in [("DeepSeek reconstruction", result["deepseek"]), ("GPT-5.6, all 60 items", result["all_60_gpt56"])] \
            + [(f"GPT-5.6 {d.split('_')[0]}", r) for d, r in result["per_dimension_gpt56"].items()]:
        lines.append(f"| {name} | {r.get('rho')} | {r.get('ci95', '')} | {r.get('n_items')} |")
    lines += ["", f"Within-question order agreement (GPT-5.6 vs expert, o3-mini vs Llama): "
                  f"{result['order_agreement_gpt56']}", "", "| Source | expert agreement | GPT-5.6 | DeepSeek | "
              "plausibility | testability | usefulness |", "|---|---:|---:|---:|---:|---:|---:|"]
    for src, v in result["by_source"].items():
        lines.append(f"| {src} | {v['expert_agreement']} | {v['gpt56_recon']} | {v['deepseek_recon']} | "
                     f"{v['plausibility']} | {v['testability']} | {v['usefulness']} |")
    lines += ["", f"Expert-rated o3-mini − Llama: {json.dumps(result['expert_o3mini_minus_llama'])}", "",
              f"Answer-key accuracy: {result['answer_key_accuracy']}",
              f"Question sense (1–5): {result['question_sense']}; hints: {result['question_hints']}",
              f"Reference identification: {result['reference_identification']}", ""]
    return "\n".join(lines)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    parser.add_argument("ratings", type=Path, help="returned workbook (.xlsx) or a folder with the four CSV files")
    args = parser.parse_args()
    ratings = read_ratings(args.ratings)
    key = json.loads((INTERNAL / "key.json").read_text())
    scores = json.loads((INTERNAL / "judge_scores.json").read_text())
    result = analyse(ratings, key, scores)
    safe = "".join(c for c in result["rater"] if c.isalnum() or c in "-_") or "expert"
    atomic_json(INTERNAL / f"analysis_{safe}.json", result)
    (INTERNAL / f"analysis_{safe}.md").write_text(render(result))
    print(render(result))


if __name__ == "__main__":
    main()
